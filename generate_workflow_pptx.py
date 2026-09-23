import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_packsure_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Color Palette matching SIH standard deck in reference image
    NAVY = RGBColor(15, 32, 67)
    DEEP_BLUE = RGBColor(28, 63, 148)
    BORDER_BLUE = RGBColor(180, 205, 245)
    BORDER_CYAN = RGBColor(147, 197, 253)
    ACCENT_ORANGE = RGBColor(234, 88, 12)
    ACCENT_GREEN = RGBColor(16, 185, 129)
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(85, 100, 120)
    CARD_BG = RGBColor(255, 255, 255)
    SOFT_BLUE_BG = RGBColor(240, 246, 255)
    HEADER_RED_BG = RGBColor(254, 226, 226)
    HEADER_RED_TXT = RGBColor(185, 28, 28)

    # 1. Slide Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    # ----------------------------------------------------
    # 2. HEADER
    # ----------------------------------------------------
    # Team Oval / Rounded Box (Left)
    left_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.16), Inches(2.1), Inches(0.6))
    left_badge.fill.solid()
    left_badge.fill.fore_color.rgb = CARD_BG
    left_badge.line.color.rgb = DEEP_BLUE
    left_badge.line.width = Pt(1.5)
    tf_lb = left_badge.text_frame
    tf_lb.word_wrap = False
    tf_lb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_lb.paragraphs[0]
    p.text = "⚡ PackSure"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE
    p.alignment = PP_ALIGN.CENTER

    # Main Title (Center)
    title_box = slide.shapes.add_textbox(Inches(2.8), Inches(0.1), Inches(7.4), Inches(0.7))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_t = tf_title.paragraphs[0]
    p_t.text = "TECHNICAL APPROACH"
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY
    p_t.alignment = PP_ALIGN.CENTER
    p_sub = tf_title.add_paragraph()
    p_sub.text = "Legal Metrology (Packaged Commodities) Rules, 2011 AI Inspection Platform"
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.alignment = PP_ALIGN.CENTER

    # SIH Badge (Right)
    sih_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.16), Inches(2.4), Inches(0.6))
    sih_badge.fill.solid()
    sih_badge.fill.fore_color.rgb = CARD_BG
    sih_badge.line.color.rgb = ACCENT_ORANGE
    sih_badge.line.width = Pt(1.5)
    tf_sih = sih_badge.text_frame
    tf_sih.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_sih = tf_sih.paragraphs[0]
    p_sih.text = "SMART INDIA HACKATHON"
    p_sih.font.size = Pt(9.5)
    p_sih.font.bold = True
    p_sih.font.color.rgb = ACCENT_ORANGE
    p_sih.alignment = PP_ALIGN.CENTER
    p_sih_id = tf_sih.add_paragraph()
    p_sih_id.text = "2026 | Problem ID: SIH26034"
    p_sih_id.font.size = Pt(8)
    p_sih_id.font.bold = True
    p_sih_id.font.color.rgb = NAVY
    p_sih_id.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # 3. LEFT PANEL - 4 ROLES FROM CODEBASE (ADMIN, SENIOR OFFICER, INSPECTOR, COMPANY)
    # ----------------------------------------------------
    roles_outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.35), Inches(0.88), Inches(1.9), Inches(5.15))
    roles_outer.fill.solid()
    roles_outer.fill.fore_color.rgb = CARD_BG
    roles_outer.line.color.rgb = BORDER_BLUE
    roles_outer.line.width = Pt(1.2)

    # Ribbon Title "ROLES"
    role_header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), Inches(0.98), Inches(1.68), Inches(0.38))
    role_header.fill.solid()
    role_header.fill.fore_color.rgb = HEADER_RED_BG
    role_header.line.color.rgb = RGBColor(239, 68, 68)
    role_header.line.width = Pt(1)
    tf_rh = role_header.text_frame
    tf_rh.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_rh = tf_rh.paragraphs[0]
    p_rh.text = "R O L E S"
    p_rh.font.size = Pt(11)
    p_rh.font.bold = True
    p_rh.font.color.rgb = HEADER_RED_TXT
    p_rh.alignment = PP_ALIGN.CENTER

    project_roles = [
        "ADMIN",
        "SENIOR\nOFFICER",
        "FIELD\nINSPECTOR",
        "COMPANY /\nMANUFACTURER"
    ]

    for idx, rname in enumerate(project_roles):
        y_pos = 1.55 + (idx * 1.08)
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), Inches(y_pos), Inches(1.68), Inches(0.85))
        pill.fill.solid()
        pill.fill.fore_color.rgb = CARD_BG
        pill.line.color.rgb = BORDER_CYAN
        pill.line.width = Pt(1.5)
        tf_p = pill.text_frame
        tf_p.word_wrap = True
        tf_p.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf_p.paragraphs[0]
        p1.text = rname
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # 4. RIGHT PANEL - TECHNOLOGY STACK (REAL PROJECT STACK)
    # ----------------------------------------------------
    tech_outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.8), Inches(0.88), Inches(2.18), Inches(5.15))
    tech_outer.fill.solid()
    tech_outer.fill.fore_color.rgb = CARD_BG
    tech_outer.line.color.rgb = DEEP_BLUE
    tech_outer.line.width = Pt(1.5)

    tech_title = slide.shapes.add_textbox(Inches(10.85), Inches(0.95), Inches(2.05), Inches(0.35))
    p_tt = tech_title.text_frame.paragraphs[0]
    p_tt.text = "Technology Stack"
    p_tt.font.size = Pt(12)
    p_tt.font.bold = True
    p_tt.font.color.rgb = NAVY
    p_tt.alignment = PP_ALIGN.CENTER

    tech_categories = [
        ("Web & Mobile Client", "React 19 & TypeScript\nVite & Tailwind CSS\nHTML5 Barcode/QR Scanner"),
        ("Backend Services", "Python 3.14 & Flask REST\nSQLAlchemy ORM\nJWT Role-Based Auth"),
        ("Vision & AI Engine", "EasyOCR & PyTesseract\nOpenCV Surface Detector\nPDP Area Geometry Engine"),
        ("Database & Reports", "PostgreSQL / SQLite\nSHA-256 Audit Trail\nReportLab PDF Generator")
    ]

    for idx, (cat_name, cat_tools) in enumerate(tech_categories):
        y_pos = 1.38 + (idx * 1.13)
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.92), Inches(y_pos), Inches(1.94), Inches(1.02))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = SOFT_BLUE_BG
        sub_card.line.color.rgb = BORDER_BLUE
        sub_card.line.width = Pt(1)
        tf_sc = sub_card.text_frame
        tf_sc.word_wrap = True
        tf_sc.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf_sc.paragraphs[0]
        p1.text = cat_name
        p1.font.size = Pt(8.5)
        p1.font.bold = True
        p1.font.color.rgb = DEEP_BLUE
        p1.alignment = PP_ALIGN.CENTER
        p2 = tf_sc.add_paragraph()
        p2.text = cat_tools
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # 5. CENTER AREA - EXACT PROJECT WORKFLOW (ACTUAL CODEBASE MODULES)
    # ----------------------------------------------------
    center_outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.4), Inches(0.88), Inches(8.25), Inches(5.15))
    center_outer.fill.solid()
    center_outer.fill.fore_color.rgb = CARD_BG
    center_outer.line.color.rgb = BORDER_BLUE
    center_outer.line.width = Pt(1.2)

    # Node builder function: Title + Exactly 2 lines
    def add_process_node(left, top, width, height, title, line1, line2, border_color=DEEP_BLUE, bg_color=CARD_BG, title_color=DEEP_BLUE):
        node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        node.fill.solid()
        node.fill.fore_color.rgb = bg_color
        node.line.color.rgb = border_color
        node.line.width = Pt(1.2)
        tf = node.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.06)
        tf.margin_right = Inches(0.06)
        tf.margin_top = Inches(0.04)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(8.5)
        p0.font.bold = True
        p0.font.color.rgb = title_color
        p0.alignment = PP_ALIGN.CENTER

        p1 = tf.add_paragraph()
        p1.text = line1
        p1.font.size = Pt(7.5)
        p1.font.color.rgb = TEXT_DARK
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = line2
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER
        return node

    # Arrow builder function
    def add_arrow(left, top, width, height, direction="RIGHT"):
        shape_type = MSO_SHAPE.RIGHT_ARROW if direction == "RIGHT" else MSO_SHAPE.DOWN_ARROW
        arr = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
        arr.fill.solid()
        arr.fill.fore_color.rgb = ACCENT_ORANGE
        arr.line.fill.background()
        return arr

    # --- TOP ENCLOSURE: SENIOR OFFICER & GOVERNANCE ---
    gov_zone = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.55), Inches(0.98), Inches(3.85), Inches(2.05))
    gov_zone.fill.solid()
    gov_zone.fill.fore_color.rgb = RGBColor(255, 251, 240)
    gov_zone.line.color.rgb = RGBColor(245, 158, 11)
    gov_zone.line.width = Pt(1.2)

    gov_lbl = slide.shapes.add_textbox(Inches(2.6), Inches(1.0), Inches(3.7), Inches(0.25))
    p_gl = gov_lbl.text_frame.paragraphs[0]
    p_gl.text = "Senior Officer: Audit Governance & Scheduling"
    p_gl.font.size = Pt(8.5)
    p_gl.font.bold = True
    p_gl.font.color.rgb = RGBColor(180, 83, 9)

    # 1. Audit Scheduling (Senior Upcoming Audits View)
    add_process_node(2.65, 1.3, 1.6, 0.78, "Audit Scheduling", "Schedules plant inspection", "Selects category & target date")

    # Arrow -> Smart Eligibility
    add_arrow(4.32, 1.6, 0.22, 0.16, "RIGHT")

    # 2. Smart Eligibility Matcher (Jurisdiction & Category Eligibility Matrix)
    add_process_node(4.6, 1.3, 1.68, 0.78, "Eligibility Matching", "Checks inspector district match", "Verifies category qualification", border_color=ACCENT_GREEN, title_color=ACCENT_GREEN)

    # Arrow Down to Dispatch
    add_arrow(4.32, 2.14, 0.16, 0.16, "DOWN")

    # 3. Inspection Dispatch & Tasking
    add_process_node(3.4, 2.36, 2.05, 0.62, "Inspector Assignment", "Dispatches assigned audit task", "Generates immutable audit record")

    # --- FIELD INSPECTOR WORKFLOW (5-STEP WIZARD) ---
    add_arrow(3.4, 3.02, 0.16, 0.22, "DOWN")

    # 4. Step 1: Product Selection & Barcode Scan
    add_process_node(2.55, 3.28, 1.75, 0.78, "Step 1: Product & Barcode", "Live Barcode / QR scanning", "Auto-fills product metadata")

    # Arrow -> Step 2
    add_arrow(4.35, 3.58, 0.2, 0.16, "RIGHT")

    # 5. Step 2: Multi-Surface Evidence Capture
    add_process_node(4.6, 3.28, 1.7, 0.78, "Step 2: Multi-Surface", "Uploads PDP, Back & Sides", "Captures batch & MRP label")

    # Arrow Down to Surface AI Check
    add_arrow(4.35, 4.1, 0.16, 0.16, "DOWN")

    # 6. Vision Analyzer: Surface Mismatch Detection
    add_process_node(3.4, 4.3, 2.05, 0.65, "Surface AI Verification", "Detects inverted PDP/Back", "Validates photo clarity & angle", bg_color=SOFT_BLUE_BG)

    # --- CENTER & RIGHT STREAM: OCR, RULE ENGINE & REVIEW ---
    add_arrow(5.5, 4.54, 0.9, 0.16, "RIGHT")

    # 7. Step 3: AI OCR & Spatial Feature Parser
    add_process_node(6.5, 1.02, 1.85, 0.85, "Step 3: AI OCR Parsing", "EasyOCR & Tesseract engines", "Parses 8 mandatory declarations")

    # Arrow Down to Deterministic Rule Engine
    add_arrow(7.35, 1.92, 0.16, 0.18, "DOWN")

    # 8. Step 4: Deterministic Statutory Rule Engine (LMPC 2011)
    add_process_node(6.5, 2.14, 1.85, 0.98, "Step 4: Statutory Engine", "Rule 6, 7, 5 & 18 verification", "PDP Area vs font height check", border_color=RGBColor(99, 102, 241), bg_color=RGBColor(238, 242, 255), title_color=RGBColor(67, 56, 202))

    # Arrow Down to Review
    add_arrow(7.35, 3.16, 0.16, 0.18, "DOWN")

    # 9. Step 5 & Senior Review Workspace (Human-in-the-Loop)
    add_process_node(6.5, 3.38, 1.85, 0.85, "Step 5 & Quasi-Judicial", "Inspector override with remarks", "Senior Officer final sign-off")

    # --- STATUTORY OUTPUTS & ARTIFACTS ---
    # 10. Court-Admissible PDF Inspection Report
    add_process_node(8.65, 1.25, 1.85, 0.82, "PDF Inspection Report", "Official inspection summary", "Complete timestamped evidence", bg_color=SOFT_BLUE_BG)

    add_arrow(8.4, 1.6, 0.2, 0.16, "RIGHT")

    # 11. Statutory Show-Cause Notice (Section 18 & 39)
    add_process_node(8.65, 2.35, 1.85, 0.85, "Legal Show-Cause Notice", "Statutory violation notice", "Cites exact rule violations", border_color=RGBColor(248, 113, 113), bg_color=RGBColor(254, 242, 242), title_color=RGBColor(185, 28, 28))

    add_arrow(8.4, 2.7, 0.2, 0.16, "RIGHT")

    # 12. SHA-256 Vault & Company Remediation Portal
    add_process_node(8.65, 3.48, 1.85, 0.85, "SHA-256 Evidence Vault", "Cryptographic tamper-proofing", "Company remediation portal", border_color=ACCENT_GREEN, title_color=ACCENT_GREEN)

    add_arrow(8.4, 3.82, 0.2, 0.16, "RIGHT")

    # --- BOTTOM SUPPORT BARS ---
    # 13. Admin Master Configuration
    add_process_node(2.55, 5.15, 3.75, 0.72, "Admin Master Data Management", "Configures rules, categories & jurisdictions", "Manages inspector eligibility matrix", border_color=BORDER_BLUE, bg_color=CARD_BG, title_color=NAVY)

    # 14. Systemic Intelligence & Analytics
    add_process_node(6.5, 5.15, 4.0, 0.72, "Systemic Intelligence & Analytics", "Detects cross-plant recurring violations", "District & state violation heatmaps", border_color=BORDER_BLUE, bg_color=CARD_BG, title_color=NAVY)

    # ----------------------------------------------------
    # 6. BOTTOM ROW - 4 CORE SYSTEM CAPABILITIES (PROJECT ACCURACY)
    # ----------------------------------------------------
    bottom_features = [
        ("Deterministic LMPC Engine", "Formally checks LMPC Rules 2011 (Rule 5, 6, 7 & 18)", "Guarantees 0% AI hallucination with exact statutory citations"),
        ("Surface AI & OCR Extraction", "Detects inverted PDP/Back labels via Vision Analyzer", "Extracts 8 mandatory declarations & measures font height in mm"),
        ("Quasi-Judicial Multi-Tier Review", "Field Inspector overrides with mandatory justification", "Senior Officer adjudication, return requests & notice approval"),
        ("SHA-256 Vault & Systemic AI", "Cryptographic image hashing for court admissibility", "Cross-inspection pattern recognition across manufacturer plants")
    ]

    for idx, (ftitle, fline1, fline2) in enumerate(bottom_features):
        x_pos = 0.35 + (idx * 3.16)
        hl_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_pos), Inches(6.15), Inches(3.06), Inches(1.18))
        hl_box.fill.solid()
        hl_box.fill.fore_color.rgb = CARD_BG
        hl_box.line.color.rgb = BORDER_BLUE
        hl_box.line.width = Pt(1.2)
        tf_h = hl_box.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = Inches(0.08)
        tf_h.margin_right = Inches(0.08)
        tf_h.margin_top = Inches(0.05)
        tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        p1 = tf_h.paragraphs[0]
        p1.text = ftitle
        p1.font.size = Pt(8.8)
        p1.font.bold = True
        p1.font.color.rgb = DEEP_BLUE
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf_h.add_paragraph()
        p2.text = fline1
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = TEXT_DARK
        p2.alignment = PP_ALIGN.CENTER

        p3 = tf_h.add_paragraph()
        p3.text = fline2
        p3.font.size = Pt(7.5)
        p3.font.color.rgb = TEXT_MUTED
        p3.alignment = PP_ALIGN.CENTER

    # Save to primary file paths
    save_paths = [
        "c:\\Project\\Hackathon\\SIH\\Packsure\\PackSure_Technical_Approach_SIH.pptx",
        "c:\\Project\\Hackathon\\SIH\\Packsure\\PackSure_Technical_Approach_Workflow_v4.pptx"
    ]
    
    for pth in save_paths:
        try:
            prs.save(pth)
            print(f"Presentation saved successfully to: {pth}")
        except Exception as e:
            print(f"Could not save to {pth}: {e}")

if __name__ == "__main__":
    create_packsure_presentation()
